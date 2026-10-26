import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "暑假重点学习资料清单_截至2026-07-11.xlsx"
OUTPUT = ROOT / "data" / "learning-data.js"


CATEGORY_ORDER = [
    "“十五五”规划",
    "全国两会",
    "二十届四中全会",
    "新年贺词",
    "中央经济工作会议",
    "中央城市工作会议",
    "中央一号文件",
    "建党105周年讲话",
    "国家科技奖励大会/两院院士大会",
    "近两年《求是》",
]


def records(ws):
    headers = [cell.value for cell in ws[1]]
    for values in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(headers, values))


def slug(value):
    value = re.sub(r"\s+", "-", str(value).strip().lower())
    return re.sub(r"[^0-9a-z\u4e00-\u9fff-]", "", value)[:42]


def main():
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    resources = []

    for index, item in enumerate(records(workbook["重点专题文章"]), 1):
        resources.append({
            "id": f"article-{index:03d}-{slug(item['标题'])}",
            "category": item["专题"],
            "format": "article",
            "kind": item["资料类型"],
            "title": item["标题"],
            "source": item["发布机构/来源"],
            "date": item["发布日期"],
            "priority": item["优先级"],
            "url": item["链接"],
            "summary": item["推荐理由/阅读重点"],
            "tags": item["关联标签"] or "",
        })

    for index, item in enumerate(records(workbook["近两年《求是》"]), 1):
        resources.append({
            "id": f"qiushi-{item['年份']}-{int(item['期号']):02d}",
            "category": "近两年《求是》",
            "format": "article",
            "kind": "《求是》重要文章",
            "title": item["标题"],
            "source": item["来源"],
            "date": item["发布日期"],
            "priority": item["建议"],
            "url": item["原文链接"],
            "summary": f"{item['年份']}年第{int(item['期号'])}期，习近平署名重要文章。",
            "tags": "习近平；求是",
        })

    for index, item in enumerate(records(workbook["视频"]), 1):
        resources.append({
            "id": f"video-{index:03d}-{slug(item['标题'])}",
            "category": item["专题"],
            "format": "video",
            "kind": item["资料类型"],
            "title": item["标题"],
            "source": item["发布机构/来源"],
            "date": item["发布日期"],
            "priority": item["优先级"],
            "url": item["链接"],
            "summary": item["推荐理由/阅读重点"],
            "tags": "视频",
        })

    resources.sort(key=lambda item: (
        CATEGORY_ORDER.index(item["category"]),
        1 if item["format"] == "video" else 0,
        item["date"],
        item["title"],
    ))

    payload = {
        "generatedAt": "2026-07-11",
        "categories": CATEGORY_ORDER,
        "resources": resources,
    }
    OUTPUT.write_text(
        "window.LEARNING_DATA = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT.name}: {len(resources)} resources")


if __name__ == "__main__":
    main()
